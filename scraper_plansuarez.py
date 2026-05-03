#!/usr/bin/env python3
"""
SCRAPER PLAN SUAREZ - VERSIÓN AUTOMÁTICA Y ESCALABLE
- Sin interfaz gráfica
- Guarda CSV y Excel localmente
- Sube directamente a Supabase
- Asigna producto_referencia_id consultando productos propios en BD
"""

import time
import csv
import re
import argparse
from datetime import datetime

import requests
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from bs4 import BeautifulSoup

from database2 import SessionLocal, PrecioHistorico, ProductoReferencia

# ============================================================================
# CONFIGURACIÓN
# ============================================================================
CATEGORIAS = {
    "Arroz": {"url": "https://www.plansuarez.com/index.php?route=product/category&path=202", "nombre": "Arroz"},
    "Harina_Maiz_Precocida": {"url": "https://www.plansuarez.com/index.php?route=product/category&path=213", "nombre": "Harina de Maíz Precocida"},
    "Harina_Trigo": {"url": "https://www.plansuarez.com/index.php?route=product/category&path=214", "nombre": "Harina de Trigo"},
    "Mayonesas": {"url": "https://www.plansuarez.com/index.php?route=product/category&path=218", "nombre": "Mayonesas"},
    "Granos": {"url": "https://www.plansuarez.com/index.php?route=product/category&path=211", "nombre": "Granos"},
    "Aceites": {"url": "https://www.plansuarez.com/index.php?route=product/category&path=201", "nombre": "Aceites"},
    "Embutidos_y_Salchichas": {"url": "https://www.plansuarez.com/index.php?route=product/category&path=154_159", "nombre": "Embutidos y Salchichas"},
    "Aves_Carnes_Cerdos": {"url": "https://www.plansuarez.com/index.php?route=product/category&path=154_158", "nombre": "Aves, Carnes y Cerdos"},
    "Ingredientes_Para_Postres": {"url": "https://www.plansuarez.com/index.php?route=product/category&path=238", "nombre": "Ingredientes para Postres"},
    "Mezclas_Pastas_Harina": {"url": "https://www.plansuarez.com/index.php?route=product/category&path=240", "nombre": "Mezclas para Pastas y Harinas"},
}

USER_AGENT = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36"
TIMEOUT = 20
DELAY_BETWEEN_PAGES = 2
OUTPUT_CSV = "Historico_PlanSuarez.csv"
OUTPUT_EXCEL = "Historico_PlanSuarez.xlsx"
SUPERMERCADO_ID = 2  # ID de Plan Suárez en Supabase

MARCAS_PROPIAS = ['La Lucha', 'Punta de Monte', 'Alibal', 'Purolomo', 'San Blas', 'Purovo', 'Milpa']

# ============================================================================
# FUNCIONES AUXILIARES
# ============================================================================
def obtener_tasa_bcv():
    try:
        res = requests.get("https://www.bcv.org.ve/", verify=False, timeout=10)
        soup = BeautifulSoup(res.text, 'html.parser')
        dolar_div = soup.find("div", {"id": "dolar"})
        if dolar_div:
            strong = dolar_div.find("strong")
            if strong:
                texto = strong.text.strip()
                match = re.search(r'(\d{2,3},\d{2,})', texto)
                if match:
                    return float(match.group(1).replace(',', '.'))
        strongs = soup.find_all("strong")
        for s in strongs:
            texto = s.text.strip()
            match = re.search(r'(\d{2,3},\d{2,})', texto)
            if match:
                return float(match.group(1).replace(',', '.'))
        return 55.0
    except Exception:
        return 55.0

def log(mensaje):
    print(f"[{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}] {mensaje}")

# ============================================================================
# CLASE SCRAPER
# ============================================================================
class PlanSuarezScraperAutomatico:
    def __init__(self, headless=False, categorias_seleccionadas=None):
        self.headless = headless
        self.categorias_a_procesar = categorias_seleccionadas or list(CATEGORIAS.keys())
        self.driver = None
        self.productos = []
        self.productos_supabase = []
        self.tasa_bcv = obtener_tasa_bcv()
        self.fecha_actual = datetime.now().date()
        self.id_contador = 1
        self.productos_propios = []

    def cargar_productos_propios(self):
        session = SessionLocal()
        try:
            productos = session.query(ProductoReferencia).filter(
                ProductoReferencia.marca.in_(MARCAS_PROPIAS),
                ProductoReferencia.activo == True
            ).all()
            self.productos_propios = [(p.nombre_producto.lower(), p.id) for p in productos]
            log(f"Cargados {len(self.productos_propios)} productos propios desde Supabase.")
        except Exception as e:
            log(f"Error cargando productos propios: {e}")
            self.productos_propios = []
        finally:
            session.close()

    def asignar_id_propio(self, nombre_original):
        nombre_norm = nombre_original.lower()
        for nombre_propio, pid in self.productos_propios:
            if nombre_propio in nombre_norm:
                return pid
        return None

    def setup_driver(self):
        options = webdriver.ChromeOptions()
        if self.headless:
            options.add_argument("--headless")
        options.add_argument(f"user-agent={USER_AGENT}")
        options.add_argument("--disable-blink-features=AutomationControlled")
        options.add_experimental_option("excludeSwitches", ["enable-automation"])
        options.add_argument("--no-sandbox")
        options.add_argument("--disable-dev-shm-usage")
        options.add_argument("--disable-gpu")
        options.add_argument("--window-size=1920,1080")
        self.driver = webdriver.Chrome(options=options)
        self.driver.execute_script("Object.defineProperty(navigator, 'webdriver', {get: () => undefined})")
        log("Driver configurado")

    def obtener_total_paginas(self, url_base):
        self.driver.get(url_base)
        time.sleep(5)
        try:
            paginacion = self.driver.find_elements(By.CSS_SELECTOR, ".pagination li a")
            if paginacion:
                ultimo = paginacion[-1]
                href = ultimo.get_attribute('href')
                if href and 'page=' in href:
                    match = re.search(r'page=(\d+)', href)
                    if match:
                        return int(match.group(1))
        except:
            pass
        return 1

    def extraer_productos_pagina(self, url_pagina):
        self.driver.get(url_pagina)
        time.sleep(DELAY_BETWEEN_PAGES)
        productos = []
        try:
            WebDriverWait(self.driver, TIMEOUT).until(
                EC.presence_of_element_located((By.CSS_SELECTOR, "div.product-layout"))
            )
            contenedores = self.driver.find_elements(By.CSS_SELECTOR, "div.product-layout")
            for cont in contenedores:
                try:
                    nombre = cont.find_element(By.CSS_SELECTOR, "div.name a").text.strip()
                except:
                    continue
                try:
                    precio_elem = cont.find_element(By.CSS_SELECTOR, "span.price-normal")
                    precio_texto = precio_elem.text.strip()
                    if not precio_texto:
                        continue
                    # Limpiar precio
                    precio_texto = precio_texto.replace("Bs.", "").replace("Bs", "").strip()
                    if ',' in precio_texto and '.' in precio_texto:
                        if precio_texto.rfind('.') > precio_texto.rfind(','):
                            precio_texto = precio_texto.replace(',', '')
                        else:
                            precio_texto = precio_texto.replace('.', '').replace(',', '.')
                    elif ',' in precio_texto:
                        precio_texto = precio_texto.replace(',', '.')
                    precio_bs = float(precio_texto)
                    precio_usd = precio_bs / self.tasa_bcv
                    productos.append((nombre, f"{precio_bs:.2f}", f"{precio_usd:.2f}"))
                except:
                    continue
        except Exception as e:
            log(f"Error extrayendo página {url_pagina}: {e}")
        return productos

    def procesar_categoria(self, categoria_key, categoria_info):
        log(f"Procesando categoría: {categoria_info['nombre']}")
        url_base = categoria_info["url"]
        total_paginas = self.obtener_total_paginas(url_base)
        log(f"Total páginas: {total_paginas}")
        for pagina in range(1, total_paginas + 1):
            url_pagina = url_base if pagina == 1 else f"{url_base}&page={pagina}"
            productos_pagina = self.extraer_productos_pagina(url_pagina)
            if not productos_pagina:
                continue
            log(f"  Página {pagina}: {len(productos_pagina)} productos")
            for nombre, precio_bs, precio_usd in productos_pagina:
                ref_id = self.asignar_id_propio(nombre)
                self.productos.append({
                    "id": self.id_contador,
                    "fecha": self.fecha_actual.strftime("%Y-%m-%d"),
                    "categoria_principal": categoria_key,
                    "nombre": nombre,
                    "precio_bs": precio_bs,
                    "precio_usd": precio_usd,
                    "tasa_bcv_usd": f"{self.tasa_bcv:.2f}",
                    "precio_ref": "",
                    "pagina": pagina,
                    "url_pagina": url_pagina,
                    "timestamp": datetime.now().isoformat()
                })
                self.id_contador += 1
                self.productos_supabase.append({
                    "supermercado_id": SUPERMERCADO_ID,
                    "nombre_original": nombre,
                    "precio_bs": precio_bs,
                    "precio_usd": precio_usd,
                    "tasa_bcv": self.tasa_bcv,
                    "fecha_extraccion": self.fecha_actual,
                    "url_pagina": url_pagina,
                    "categoria_extraida": categoria_key,
                    "producto_referencia_id": ref_id,
                    "matched_automatically": ref_id is not None,
                    "match_score": 100 if ref_id else None
                })

    def guardar_csv_excel(self):
        if not self.productos:
            log("No hay productos para guardar localmente.")
            return
        with open(OUTPUT_CSV, 'w', newline='', encoding='utf-8-sig') as f:
            fieldnames = ['id', 'fecha', 'categoria_principal', 'nombre', 'precio_bs',
                          'precio_usd', 'tasa_bcv_usd', 'precio_ref', 'pagina', 'url_pagina', 'timestamp']
            writer = csv.DictWriter(f, fieldnames=fieldnames)
            writer.writeheader()
            writer.writerows(self.productos)
        log(f"CSV guardado: {OUTPUT_CSV} ({len(self.productos)} registros)")
        try:
            import openpyxl
            from openpyxl.styles import PatternFill, Font, Alignment
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Historico_PlanSuarez"
            with open(OUTPUT_CSV, 'r', encoding='utf-8-sig') as f:
                reader = csv.reader(f)
                for r_idx, row in enumerate(reader, 1):
                    for c_idx, val in enumerate(row, 1):
                        ws.cell(row=r_idx, column=c_idx, value=val)
            header_fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
            header_font = Font(color="FFFFFF", bold=True)
            for col in range(1, ws.max_column + 1):
                cell = ws.cell(row=1, column=col)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center")
            for col in range(1, ws.max_column + 1):
                ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 25
            wb.save(OUTPUT_EXCEL)
            log(f"Excel guardado: {OUTPUT_EXCEL}")
        except Exception as e:
            log(f"No se pudo crear Excel: {e}")

    def subir_a_supabase(self):
        if not self.productos_supabase:
            log("No hay datos para subir.")
            return
        session = SessionLocal()
        try:
            session.bulk_insert_mappings(PrecioHistorico, self.productos_supabase)
            session.commit()
            log(f"✅ Subidos {len(self.productos_supabase)} registros a Supabase (Plan Suárez).")
        except Exception as e:
            session.rollback()
            log(f"❌ Error en Supabase: {e}")
        finally:
            session.close()

    def ejecutar(self):
        log("="*60)
        log("INICIANDO SCRAPER PLAN SUAREZ (AUTOMÁTICO)")
        log(f"Tasa BCV: {self.tasa_bcv:.2f}")
        self.cargar_productos_propios()
        if not self.productos_propios:
            log("⚠️ No se encontraron productos propios. Se asignará NULL.")
        log(f"Categorías: {', '.join(self.categorias_a_procesar)}")
        try:
            self.setup_driver()
            for cat_key in self.categorias_a_procesar:
                if cat_key not in CATEGORIAS:
                    continue
                self.procesar_categoria(cat_key, CATEGORIAS[cat_key])
                time.sleep(2)
            self.guardar_csv_excel()
            self.subir_a_supabase()
        except Exception as e:
            log(f"Error crítico: {e}")
        finally:
            if self.driver:
                self.driver.quit()
        log("SCRAPER FINALIZADO")

if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--headless", action="store_true", help="Ejecutar sin ventana")
    parser.add_argument("--categorias", nargs="+", help="Categorías específicas")
    args = parser.parse_args()
    scraper = PlanSuarezScraperAutomatico(headless=args.headless, categorias_seleccionadas=args.categorias)
    scraper.ejecutar()