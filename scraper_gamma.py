#!/usr/bin/env python3
"""
SCRAPER GAMAENLINEA - VERSIÓN AUTOMÁTICA Y ESCALABLE
- Sin interfaz gráfica
- Guarda CSV y Excel localmente
- Sube los datos a Supabase
- Asigna producto_referencia_id consultando los productos propios en Supabase
- No requiere MAPEO manual: usa los nombres de productos propios de la BD
"""

import time
import csv
import re
import argparse
from datetime import datetime

import requests
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from bs4 import BeautifulSoup

from database2 import SessionLocal, PrecioHistorico, ProductoReferencia

# ============================================================================
# CONFIGURACIÓN
# ============================================================================
CATEGORIAS = {
    "Harinas_y_Pastas": {"url": "https://gamaenlinea.com/es/harinas-pastas/c/A0102", "nombre": "Harinas y Pastas"},
    "Arroz_y_Granos": {"url": "https://gamaenlinea.com/es/arroz-granos/c/A0103", "nombre": "Arroz y Granos"},
    "Aceites_y_Vinagres": {"url": "https://gamaenlinea.com/es/aceites-vinagres/c/A0109", "nombre": "Aceites y Vinagres"},
    "Mezclas_Postres": {"url": "https://gamaenlinea.com/es/mezclas-postres-e-insumos/c/A010604", "nombre": "Mezclas para Postres e Insumos"},
    "Mayonesa": {"url": "https://gamaenlinea.com/es/mayonesa/c/A011005", "nombre": "Mayonesa"},
    "Cereales": {"url": "https://gamaenlinea.com/es/cereales/c/A0115", "nombre": "Cereales"},
    "Charcuteria": {"url": "https://gamaenlinea.com/es/charcuteria/c/A0202", "nombre": "Charcutería"},
    "Huevos": {"url": "https://gamaenlinea.com/es/huevos/c/A0207", "nombre": "Huevos"},
    "Carniceria": {"url": "https://gamaenlinea.com/es/carniceria/c/A0201", "nombre": "Carnicería"}
}

USER_AGENT = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36"
TIMEOUT = 20
DELAY_BETWEEN_PAGES = 2
OUTPUT_CSV = "Historico_Gamma.csv"
OUTPUT_EXCEL = "Historico_Gamma.xlsx"
SUPERMERCADO_ID = 1

# Marcas que consideramos propias (para filtrar productos en la BD)
MARCAS_PROPIAS = ['La Lucha', 'Punta de Monte', 'Alibal', 'Purolomo', 'San Blas', 'Purovo', 'Milpa']

# ============================================================================
# FUNCIONES AUXILIARES
# ============================================================================
def obtener_tasa_bcv():
    try:
        res = requests.get("https://www.bcv.org.ve/", verify=False, timeout=10)
        soup = BeautifulSoup(res.text, 'html.parser')
        dolar_div = soup.find("div", {"id": "dolar"})
        if dolar_div:
            strong = dolar_div.find("strong")
            if strong:
                texto = strong.text.strip()
                match = re.search(r'(\d{2,3},\d{2,})', texto)
                if match:
                    return float(match.group(1).replace(',', '.'))
        strongs = soup.find_all("strong")
        for s in strongs:
            texto = s.text.strip()
            match = re.search(r'(\d{2,3},\d{2,})', texto)
            if match:
                return float(match.group(1).replace(',', '.'))
        return 55.0
    except:
        return 55.0

def log(mensaje):
    print(f"[{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}] {mensaje}")

# ============================================================================
# CLASE SCRAPER
# ============================================================================
class GamaScraperAutomatico:
    def __init__(self, headless=False, categorias_seleccionadas=None):
        self.headless = headless
        self.categorias_a_procesar = categorias_seleccionadas or list(CATEGORIAS.keys())
        self.driver = None
        self.productos = []          # para CSV/Excel
        self.productos_supabase = [] # para insertar
        self.tasa_bcv = obtener_tasa_bcv()
        self.fecha_actual = datetime.now().date()
        self.id_contador = 1
        self.productos_propios = []  # lista de (nombre_producto, id)

    def cargar_productos_propios(self):
        """Carga los productos propios desde la base de datos (activos y con marca propia)"""
        session = SessionLocal()
        try:
            productos = session.query(ProductoReferencia).filter(
                ProductoReferencia.marca.in_(MARCAS_PROPIAS),
                ProductoReferencia.activo == True
            ).all()
            self.productos_propios = [(p.nombre_producto.lower(), p.id) for p in productos]
            log(f"Cargados {len(self.productos_propios)} productos propios desde Supabase.")
        except Exception as e:
            log(f"Error cargando productos propios: {e}")
            self.productos_propios = []
        finally:
            session.close()

    def asignar_id_propio(self, nombre_original):
        """Devuelve el producto_referencia_id si el nombre_original coincide con algún producto propio, None en caso contrario"""
        nombre_norm = nombre_original.lower()
        # Intentar coincidencia exacta o por subcadena (el nombre propio contenido en el original)
        for nombre_propio, pid in self.productos_propios:
            if nombre_propio in nombre_norm:
                return pid
        return None

    def setup_driver(self):
        options = webdriver.ChromeOptions()
        if self.headless:
            options.add_argument("--headless")
        options.add_argument(f"user-agent={USER_AGENT}")
        options.add_argument("--disable-blink-features=AutomationControlled")
        options.add_experimental_option("excludeSwitches", ["enable-automation"])
        options.add_argument("--no-sandbox")
        options.add_argument("--disable-dev-shm-usage")
        options.add_argument("--disable-gpu")
        options.add_argument("--window-size=1920,1080")
        self.driver = webdriver.Chrome(options=options)
        self.driver.execute_script("Object.defineProperty(navigator, 'webdriver', {get: () => undefined})")
        log("Driver configurado")

    def obtener_total_paginas(self, url_base):
        self.driver.get(url_base)
        time.sleep(5)
        try:
            paginacion = self.driver.find_elements(By.CSS_SELECTOR, "eg-pagination a")
            ultimo_numero = 0
            for elem in paginacion:
                href = elem.get_attribute('href')
                if href and 'currentPage=' in href:
                    match = re.search(r'currentPage=(\d+)', href)
                    if match:
                        num = int(match.group(1)) + 1
                        if num > ultimo_numero:
                            ultimo_numero = num
            if ultimo_numero > 0:
                return ultimo_numero
        except:
            pass
        try:
            ultimo_enlace = self.driver.find_element(By.CSS_SELECTOR, "eg-pagination a.end")
            href = ultimo_enlace.get_attribute('href')
            if href and 'currentPage=' in href:
                match = re.search(r'currentPage=(\d+)', href)
                if match:
                    return int(match.group(1)) + 1
        except:
            pass
        return 1

    def extraer_productos_pagina(self, url_pagina):
        self.driver.get(url_pagina)
        time.sleep(DELAY_BETWEEN_PAGES)
        productos = []
        try:
            WebDriverWait(self.driver, TIMEOUT).until(
                EC.presence_of_element_located((By.CSS_SELECTOR, "a.cx-product-name h3"))
            )
            contenedores = self.driver.find_elements(By.CSS_SELECTOR, "cx-product-grid-item")
            for cont in contenedores:
                try:
                    nombre = cont.find_element(By.CSS_SELECTOR, "a.cx-product-name h3").text.strip()
                except:
                    continue
                precio_ref = ""
                try:
                    spans = cont.find_elements(By.CSS_SELECTOR, "div.cx-product-price span")
                    for span in spans:
                        texto = span.text.strip()
                        if "Total Ref." in texto:
                            match = re.search(r'([\d]+,[\d]{2})', texto)
                            if match:
                                precio_ref = match.group(1).replace(',', '.')
                                break
                except:
                    pass
                precio_bs = ""
                try:
                    spans = cont.find_elements(By.CSS_SELECTOR, "div.cx-product-price span")
                    for span in spans:
                        texto = span.text.strip()
                        if "Total Bs." in texto:
                            match = re.search(r'([\d\.]+,\d{2})', texto)
                            if match:
                                precio_bs = match.group(1).replace('.', '').replace(',', '.')
                                break
                except:
                    pass
                if not precio_bs and precio_ref and self.tasa_bcv > 0:
                    try:
                        precio_bs = f"{float(precio_ref) * self.tasa_bcv:.2f}"
                    except:
                        pass
                if not precio_bs:
                    continue
                try:
                    precio_usd = f"{float(precio_bs) / self.tasa_bcv:.2f}"
                except:
                    precio_usd = "0.00"
                productos.append((nombre, precio_ref, precio_bs, precio_usd))
        except Exception as e:
            log(f"Error en página {url_pagina}: {e}")
        return productos

    def procesar_categoria(self, categoria_key, categoria_info):
        log(f"Procesando categoría: {categoria_info['nombre']}")
        url_base = categoria_info["url"]
        total_paginas = self.obtener_total_paginas(url_base)
        log(f"Total páginas: {total_paginas}")
        for pagina in range(1, total_paginas + 1):
            url_pagina = url_base if pagina == 1 else f"{url_base}?currentPage={pagina-1}"
            productos_pagina = self.extraer_productos_pagina(url_pagina)
            if not productos_pagina:
                continue
            log(f"  Página {pagina}: {len(productos_pagina)} productos")
            for nombre, precio_ref, precio_bs, precio_usd in productos_pagina:
                # Asignar ID propio automáticamente
                ref_id = self.asignar_id_propio(nombre)
                # Para CSV/Excel
                self.productos.append({
                    "id": self.id_contador,
                    "fecha": self.fecha_actual.strftime("%Y-%m-%d"),
                    "categoria_principal": categoria_key,
                    "nombre": nombre,
                    "precio_bs": precio_bs,
                    "precio_usd": precio_usd,
                    "tasa_bcv_usd": f"{self.tasa_bcv:.2f}",
                    "precio_ref": precio_ref,
                    "pagina": pagina,
                    "url_pagina": url_pagina,
                    "timestamp": datetime.now().isoformat()
                })
                self.id_contador += 1
                # Para Supabase
                self.productos_supabase.append({
                    "supermercado_id": SUPERMERCADO_ID,
                    "nombre_original": nombre,
                    "precio_bs": precio_bs,
                    "precio_usd": precio_usd,
                    "tasa_bcv": self.tasa_bcv,
                    "fecha_extraccion": self.fecha_actual,
                    "url_pagina": url_pagina,
                    "categoria_extraida": categoria_key,
                    "producto_referencia_id": ref_id,
                    "matched_automatically": ref_id is not None,
                    "match_score": 100 if ref_id else None
                })

    def guardar_csv_excel(self):
        if not self.productos:
            log("No hay productos para guardar localmente.")
            return
        with open(OUTPUT_CSV, 'w', newline='', encoding='utf-8-sig') as f:
            fieldnames = ['id', 'fecha', 'categoria_principal', 'nombre', 'precio_bs',
                          'precio_usd', 'tasa_bcv_usd', 'precio_ref', 'pagina', 'url_pagina', 'timestamp']
            writer = csv.DictWriter(f, fieldnames=fieldnames)
            writer.writeheader()
            writer.writerows(self.productos)
        log(f"CSV guardado: {OUTPUT_CSV} ({len(self.productos)} registros)")
        try:
            import openpyxl
            from openpyxl.styles import PatternFill, Font, Alignment
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Historico_Gamma"
            with open(OUTPUT_CSV, 'r', encoding='utf-8-sig') as f:
                reader = csv.reader(f)
                for r_idx, row in enumerate(reader, 1):
                    for c_idx, val in enumerate(row, 1):
                        ws.cell(row=r_idx, column=c_idx, value=val)
            header_fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
            header_font = Font(color="FFFFFF", bold=True)
            for col in range(1, ws.max_column + 1):
                cell = ws.cell(row=1, column=col)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center")
            for col in range(1, ws.max_column + 1):
                ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 25
            wb.save(OUTPUT_EXCEL)
            log(f"Excel guardado: {OUTPUT_EXCEL}")
        except Exception as e:
            log(f"No se pudo crear Excel: {e}")

    def subir_a_supabase(self):
        if not self.productos_supabase:
            log("No hay datos para subir a Supabase.")
            return
        session = SessionLocal()
        try:
            session.bulk_insert_mappings(PrecioHistorico, self.productos_supabase)
            session.commit()
            log(f"✅ Subidos {len(self.productos_supabase)} registros a Supabase (Gamma).")
        except Exception as e:
            session.rollback()
            log(f"❌ Error en Supabase: {e}")
        finally:
            session.close()

    def ejecutar(self):
        log("="*60)
        log("INICIANDO SCRAPER GAMA (AUTOMÁTICO)")
        log(f"Tasa BCV: {self.tasa_bcv:.2f}")
        # Cargar productos propios desde Supabase
        self.cargar_productos_propios()
        if not self.productos_propios:
            log("⚠️ No se encontraron productos propios en la base de datos. Se asignará NULL a todos.")
        log(f"Categorías a procesar: {', '.join(self.categorias_a_procesar)}")
        try:
            self.setup_driver()
            for cat_key in self.categorias_a_procesar:
                if cat_key not in CATEGORIAS:
                    log(f"Categoría '{cat_key}' no válida, omitida.")
                    continue
                self.procesar_categoria(cat_key, CATEGORIAS[cat_key])
                time.sleep(2)
            self.guardar_csv_excel()
            self.subir_a_supabase()
        except Exception as e:
            log(f"Error crítico: {e}")
            import traceback
            traceback.print_exc()
        finally:
            if self.driver:
                self.driver.quit()
        log("SCRAPER FINALIZADO")

if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--headless", action="store_true", help="Ejecutar sin ventana del navegador")
    parser.add_argument("--categorias", nargs="+", help="Lista de categorías a extraer (ej. Harinas_y_Pastas Arroz_y_Granos)")
    args = parser.parse_args()
    scraper = GamaScraperAutomatico(headless=args.headless, categorias_seleccionadas=args.categorias)
    scraper.ejecutar()